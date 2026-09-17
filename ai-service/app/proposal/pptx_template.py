"""TTN proposal PPTX template rendering.

The template at templates/proposal-template.pptx is the base deck for every RFP.
Only the Why TTN? slide and Success Stories slides keep fixed TTN marketing content.
Every other slide is filled or rewritten from generated proposal content and customer metadata.
"""

import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any, Iterator

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

from app.config import settings

# 0-based slide indices that must never be modified (TTN boilerplate only).
WHY_TTN_SLIDE_INDEX = 3
SUCCESS_STORY_SLIDE_INDICES: frozenset[int] = frozenset({22, 23, 24})
FIXED_SLIDE_INDICES: frozenset[int] = frozenset({WHY_TTN_SLIDE_INDEX}) | SUCCESS_STORY_SLIDE_INDICES

# Duplicate agenda dividers in the source deck (still rewritten with customer-specific text).
AGENDA_SLIDE_INDICES: frozenset[int] = frozenset({1, 10, 15, 19, 21})

IN_SCOPE_SLIDE_INDEX = 9
ENGAGEMENT_SLIDE_INDEX = 20

# Legacy export: slides removed when present in older templates.
SLIDES_TO_REMOVE: frozenset[int] = frozenset()

# Sample client/project strings from the source TTN deck — replaced on dynamic slides.
_TEMPLATE_CLIENT_MARKERS: tuple[str, ...] = (
    "Racing Queensland Networks",
    "Racing Queensland",
    "Races & Profiles",
    "Mobile App Development | Racing Queensland",
)

# Phrases from the sample mobile-app deck — cleared when still present on dynamic slides.
_TEMPLATE_SAMPLE_PHRASES: tuple[str, ...] = (
    "Races & Profiles",
    "Blackbook",
    "Jockey/Trainer",
    "Race Calendar & Planner",
    "Betting Partners",
    "Queensland Non-TAB",
    "After a detailed discovery of the platform and user needs",
    "Both Flutter and React Native allow",
    "Key Application Modules",
)

MULTI_CARD_SLIDE_INDICES: frozenset[int] = frozenset({5})

# Human-readable titles when the template leaves placeholders empty.
SLIDE_TITLES: dict[int, str] = {
    2: "Executive Summary",
    4: "Our Understanding of Requirements",
    5: "Key Application Modules",
    6: "Non Functional Requirements",
    7: "Non Functional Requirements",
    8: "Key Assumptions and Dependencies",
    9: "In Scope and Out of Scope",
    11: "Solutioning Highlights",
    12: "Proposed High Level Architecture",
    13: "Architecture Considerations",
    14: "Mobile Development Platform - Comparison",
    16: "Project Management Approach",
    17: "Project Governance Model",
    18: "Project Plan",
    20: "Proposed Engagement Model",
    25: "Quick Recap",
}

ARCHITECTURE_SLIDE_INDEX = 12
GOVERNANCE_SLIDE_INDEX = 17
PROJECT_PLAN_SLIDE_INDEX = 18

# Slides with dense diagrams — short annotations only (not architecture/governance/plan).
DIAGRAM_OVERLAY_SLIDE_INDICES: frozenset[int] = frozenset({16})

# Small annotation boxes on diagram slides (content key).
DIAGRAM_ANNOTATION_CONTENT: dict[int, str] = {
    16: "implementationMethodology",
}

# Two-column tables filled from "left | right" lines.
TABLE_TWO_COLUMN_SLIDES: dict[int, str] = {13: "technicalApproach"}

# Three-column comparison table — only rows with "a | b | c" lines.
TABLE_COMPARISON_SLIDE_INDEX = 14

# 0-based slide index -> proposal content key (body / card layouts).
SLIDE_CONTENT_MAP: dict[int, str] = {
    2: "executiveSummary",
    4: "understandingOfRequirements",
    5: "proposedSolution",
    6: "securityCompliance",
    7: "technicalApproach",
    8: "assumptions",
    11: "proposedSolution",
    25: "executiveSummary",
}

# Extra proposal keys tried when the primary slide field is empty.
SLIDE_CONTENT_FALLBACKS: dict[int, tuple[str, ...]] = {
    2: ("understandingOfRequirements", "proposedSolution"),
    5: ("understandingOfRequirements", "technicalApproach"),
    6: ("technicalApproach", "supportAndSla"),
    7: ("proposedSolution", "architectureOverview"),
    8: ("understandingOfRequirements", "technicalApproach"),
    11: ("technicalApproach", "implementationMethodology"),
    13: ("architectureOverview", "proposedSolution"),
    25: ("understandingOfRequirements", "proposedSolution"),
}

EMPTY_TRAILING_SLIDE_INDEX = 26

TITLE_PLACEHOLDER_MARKERS = (
    "Click to add title",
    "Click to add subtitle",
    "Click to add text",
)

SKIP_TEXT = frozenset({"‹#›", "Commercial Summary", "Success Stories", "Agenda"})
FOOTER_TOP_EMU = 4_000_000
FOOTNOTE_TOP_EMU = 4_000_000
CHARS_PER_INCH_WIDTH = 11
LINES_PER_INCH_HEIGHT = 4.5
BODY_FONT_PT = 10
MIN_SHAPE_WIDTH_IN = 0.8
MIN_SHAPE_HEIGHT_IN = 0.5
EMU_PER_INCH = 914_400


def resolve_pptx_template_path() -> Path:
    configured = Path(settings.pptx_template_path)
    if configured.exists():
        return configured

    templates_dir = Path(__file__).resolve().parents[3] / "templates"
    for name in (
        "proposal-template.pptx",
        "TTN Proposal for Mobile App Development .pptx",
    ):
        candidate = templates_dir / name
        if candidate.exists():
            return candidate

    return configured


def _coerce_text(value: Any, max_len: int | None = None) -> str:
    if value is None:
        text = ""
    elif isinstance(value, str):
        text = value
    elif isinstance(value, list):
        text = "\n".join(_coerce_text(item) for item in value if item)
    elif isinstance(value, dict):
        for key in ("text", "content", "body", "paragraph", "summary", "value"):
            if key in value:
                text = _coerce_text(value[key])
                break
        else:
            text = "\n".join(_coerce_text(item) for item in value.values() if item)
    else:
        text = str(value)

    return text[:max_len] if max_len is not None else text


def _iter_text_shapes(slide: Any) -> Iterator[Any]:
    for shape in slide.shapes:
        try:
            if hasattr(shape, "text_frame"):
                yield shape
            elif hasattr(shape, "shapes"):
                for child in shape.shapes:
                    if hasattr(child, "text_frame"):
                        yield child
        except AttributeError:
            continue


def _shape_area(shape: Any) -> int:
    return int(getattr(shape, "width", 0)) * int(getattr(shape, "height", 0))


def _is_footer_shape(shape: Any) -> bool:
    text = shape.text.strip() if hasattr(shape, "text") else ""
    if text in SKIP_TEXT:
        return True

    top = getattr(shape, "top", 0)
    width_in, _ = _shape_size_inches(shape)
    return top > 4_700_000 and width_in < 1.5


def _find_title_shape(shapes: list[Any]) -> Any | None:
    for shape in shapes:
        if _is_footer_shape(shape):
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER:
            text = shape.text.strip()
            if text and text not in SKIP_TEXT:
                return shape

    titled: list[tuple[int, Any]] = []
    for shape in shapes:
        text = shape.text.strip()
        if not text or _is_footer_shape(shape):
            continue
        if getattr(shape, "top", 0) > 700_000:
            continue
        if len(text) > 120:
            continue
        titled.append((getattr(shape, "top", 0), shape))

    if titled:
        return min(titled, key=lambda item: item[0])[1]
    return None


def _ensure_slide_title(slide: Any, title: str) -> None:
    if not title.strip():
        return

    shapes = list(_iter_text_shapes(slide))
    title_shape = _find_title_shape(shapes)
    if title_shape is not None:
        existing = title_shape.text.strip()
        if (
            not existing
            or any(marker in existing for marker in TITLE_PLACEHOLDER_MARKERS)
        ):
            title_shape.text = title
        return

    for shape in sorted(shapes, key=lambda item: getattr(item, "top", 0)):
        if _is_footer_shape(shape):
            continue
        if getattr(shape, "top", 0) > 800_000:
            continue
        existing = shape.text.strip()
        if not existing or any(marker in existing for marker in TITLE_PLACEHOLDER_MARKERS):
            shape.text = title
            return


def _shape_size_inches(shape: Any) -> tuple[float, float]:
    return (
        getattr(shape, "width", 0) / EMU_PER_INCH,
        getattr(shape, "height", 0) / EMU_PER_INCH,
    )


def _shape_char_capacity(shape: Any) -> int:
    width_in, height_in = _shape_size_inches(shape)
    if width_in < MIN_SHAPE_WIDTH_IN or height_in < MIN_SHAPE_HEIGHT_IN:
        return 0

    chars_per_line = max(12, int(width_in * CHARS_PER_INCH_WIDTH))
    max_lines = max(2, int(height_in * LINES_PER_INCH_HEIGHT))
    return chars_per_line * max_lines


def _truncate_text(text: str, max_chars: int) -> str:
    cleaned = re.sub(r"\s+", " ", text).strip()
    if max_chars <= 0:
        return ""
    if len(cleaned) <= max_chars:
        return cleaned

    trimmed = cleaned[:max_chars].rstrip()
    for separator in (". ", "; ", ", ", " "):
        split_at = trimmed.rfind(separator)
        if split_at >= int(max_chars * 0.55):
            return trimmed[: split_at + (1 if separator == ". " else 0)].strip()

    return trimmed.rstrip(" ,;.") + "…"


def _apply_body_font(text_frame: Any) -> None:
    text_frame.word_wrap = True
    for paragraph in text_frame.paragraphs:
        paragraph.space_after = Pt(2)
        for run in paragraph.runs:
            run.font.size = Pt(BODY_FONT_PT)


def _set_text_frame(
    text_frame: Any,
    text: str,
    shape: Any | None = None,
    max_bullets: int = 12,
) -> None:
    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
    if not paragraphs:
        return

    capacity = _shape_char_capacity(shape) if shape is not None else 800
    if capacity <= 0:
        capacity = 120

    fitted: list[str] = []
    used = 0
    for paragraph in paragraphs[:max_bullets]:
        remaining = capacity - used
        if remaining <= 0:
            break
        line = _truncate_text(paragraph, remaining)
        if not line:
            break
        fitted.append(line)
        used += len(line) + 1

    if not fitted:
        fitted = [_truncate_text(paragraphs[0], capacity)]

    text_frame.clear()
    text_frame.text = fitted[0]
    for paragraph in fitted[1:]:
        bullet = text_frame.add_paragraph()
        bullet.text = paragraph
        bullet.level = 0

    _apply_body_font(text_frame)


def _content_to_lines(text: str) -> list[str]:
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    if len(lines) <= 1 and len(text) > 120:
        lines = [part.strip() for part in re.split(r"(?<=[.!?])\s+", text) if part.strip()]
    return lines


def _split_table_line(line: str, max_left: int = 80, max_right: int = 180) -> tuple[str, str]:
    for separator in (" | ", " - ", ": "):
        if separator in line:
            left, right = line.split(separator, 1)
            return _truncate_text(left, max_left), _truncate_text(right, max_right)
    if len(line) > max_left:
        return _truncate_text(line, max_left), _truncate_text(line, max_right)
    return line, line


def _is_footnote_shape(shape: Any) -> bool:
    top = getattr(shape, "top", 0)
    if top < 1_200_000:
        return False
    _, height_in = _shape_size_inches(shape)
    width_in, _ = _shape_size_inches(shape)
    return top > FOOTNOTE_TOP_EMU or (height_in < 0.75 and width_in > 6)


def _is_title_shape(shape: Any, title_shape: Any | None) -> bool:
    if title_shape is not None and shape is title_shape:
        return True
    if title_shape is not None:
        return (
            getattr(shape, "top", -1) == getattr(title_shape, "top", -2)
            and getattr(shape, "left", -1) == getattr(title_shape, "left", -2)
        )
    top = getattr(shape, "top", 0)
    text = shape.text.strip() if hasattr(shape, "text") else ""
    return top < 700_000 and bool(text) and len(text) < 120


def _clear_footnotes(slide: Any, title_shape: Any | None) -> None:
    for shape in _iter_text_shapes(slide):
        if _is_title_shape(shape, title_shape) or _is_footer_shape(shape):
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            continue
        if _is_footnote_shape(shape):
            shape.text_frame.clear()


def _clear_table_data_rows(table: Any) -> None:
    for row_index in range(1, len(table.rows)):
        for col_index in range(len(table.columns)):
            table.cell(row_index, col_index).text = ""


def _update_two_column_table(table: Any, text: str) -> bool:
    lines = _content_to_lines(text)
    filled = False
    for row_index in range(1, len(table.rows)):
        line_index = row_index - 1
        if line_index < len(lines):
            left, right = _split_table_line(lines[line_index])
            table.cell(row_index, 0).text = left
            if len(table.columns) > 1:
                table.cell(row_index, 1).text = right
            filled = True
        else:
            table.cell(row_index, 0).text = ""
            if len(table.columns) > 1:
                table.cell(row_index, 1).text = ""
    return filled


def _table_shapes(slide: Any) -> list[Any]:
    return [
        shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE
    ]


def _comparison_lines(proposal: dict[str, Any]) -> str:
    raw = _coerce_text(proposal.get("mobilePlatformComparison", ""))
    if raw and sum(1 for line in raw.split("\n") if line.count("|") >= 2) >= 2:
        return raw

    tech = _coerce_text(proposal.get("technicalApproach", "")).lower()
    solution = _coerce_text(proposal.get("proposedSolution", "")).lower()
    combined = f"{tech} {solution}"
    if "flutter" in combined and "react" not in combined:
        col_b, col_c = "Flutter", "React Native"
    elif "react native" in combined or "react" in combined:
        col_b, col_c = "React Native", "Flutter"
    else:
        col_b, col_c = "Recommended option", "Alternative option"

    return "\n".join(
        [
            f"Time to market | {col_b} | {col_c}",
            "Native performance | High | High",
            "Team skill fit | Based on existing engineering skills | Based on mobile specialization",
            "UI consistency | Strong component libraries | Strong widget catalog",
            "Integration with APIs | REST/GraphQL friendly | REST/GraphQL friendly",
            "Maintainability | Modular architecture | Modular architecture",
            "Recommendation | Aligns with stated RFP stack preferences | Viable alternate if skills differ",
        ]
    )


def _update_comparison_slide(slide: Any, proposal: dict[str, Any]) -> None:
    title = SLIDE_TITLES[TABLE_COMPARISON_SLIDE_INDEX]
    _ensure_slide_title(slide, title)
    text = _comparison_lines(proposal)
    tables = _table_shapes(slide)
    if tables:
        table = tables[0].table
        if not _update_three_column_table(table, text):
            _default_comparison_rows(table, proposal)
    recommendation = _coerce_text(proposal.get("mobilePlatformComparison", ""))
    if not recommendation:
        recommendation = (
            "Recommendation follows the technical approach and skills implied by the RFP requirements."
        )
    for shape in _iter_text_shapes(slide):
        if getattr(shape, "top", 0) > 4_200_000 and _shape_char_capacity(shape) >= 80:
            shape.text = _truncate_text(recommendation.split("\n")[-1], 220)
            _apply_body_font(shape.text_frame)
            break


def _default_comparison_rows(table: Any, proposal: dict[str, Any]) -> None:
    lines = _content_to_lines(_comparison_lines(proposal))
    _update_three_column_table(table, "\n".join(lines))


def _update_architecture_slide(slide: Any, proposal: dict[str, Any]) -> None:
    _ensure_slide_title(slide, SLIDE_TITLES[ARCHITECTURE_SLIDE_INDEX])
    _purge_template_sample_text(slide)

    text = _coerce_text(
        proposal.get("architectureOverview") or proposal.get("technicalApproach", "")
    )
    lines = _content_to_lines(text)
    customer = _customer_name(proposal)
    stack_label = _truncate_text(lines[0] if lines else f"{customer} Platform", 45)

    title_shape = _find_title_shape(list(_iter_text_shapes(slide)))
    for shape in _iter_text_shapes(slide):
        if _is_title_shape(shape, title_shape) or _is_footer_shape(shape):
            continue
        label = shape.text.strip()
        if label in {"Mobile Application", "Mobile  User", "Mobile User"}:
            shape.text = stack_label
        elif label == "Backend APIs" and len(lines) > 1:
            shape.text = _truncate_text(lines[1], 35)
        elif label == "Firebase" and len(lines) > 2:
            shape.text = _truncate_text(lines[2], 35)

    layer_boxes = sorted(
        [
            shape
            for shape in _iter_text_shapes(slide)
            if 100 <= _shape_char_capacity(shape) <= 280
            and 2_000_000 < getattr(shape, "left", 0) < 6_500_000
            and getattr(shape, "top", 0) > 2_000_000
        ],
        key=lambda item: item.top,
    )
    layer_lines = lines[1:5] if len(lines) > 1 else lines[:4]
    for index, shape in enumerate(layer_boxes[:2]):
        if index < len(layer_lines):
            _set_text_frame(shape.text_frame, layer_lines[index], shape, max_bullets=4)
        else:
            shape.text_frame.clear()

    _clear_center_overlay_text(slide, title_shape)


def _governance_notes_shape(slide: Any, title_shape: Any | None) -> Any | None:
    candidates: list[tuple[int, Any]] = []
    for shape in _iter_text_shapes(slide):
        if _is_title_shape(shape, title_shape) or _is_footer_shape(shape):
            continue
        top = getattr(shape, "top", 0)
        width = getattr(shape, "width", 0)
        capacity = _shape_char_capacity(shape)
        if 3_200_000 <= top <= 3_700_000 and 2_000_000 <= width <= 5_500_000:
            if 200 <= capacity <= 400:
                candidates.append((capacity, shape))
    if not candidates:
        return None
    return max(candidates, key=lambda item: item[0])[1]


def _clear_governance_overlays(slide: Any, title_shape: Any | None) -> None:
    for shape in _iter_text_shapes(slide):
        if _is_title_shape(shape, title_shape) or _is_footer_shape(shape):
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            continue
        width = getattr(shape, "width", 0)
        capacity = _shape_char_capacity(shape)
        if width > 4_000_000 and capacity > 80:
            shape.text_frame.clear()
        elif _shape_has_sample_content(shape) and capacity > 80:
            shape.text_frame.clear()


def _update_governance_slide(slide: Any, proposal: dict[str, Any]) -> None:
    _ensure_slide_title(slide, SLIDE_TITLES[GOVERNANCE_SLIDE_INDEX])
    title_shape = _find_title_shape(list(_iter_text_shapes(slide)))
    _clear_governance_overlays(slide, title_shape)
    _apply_customer_substitutions(slide, proposal)

    text = _coerce_text(
        proposal.get("governanceModel")
        or proposal.get("supportAndSla", "")
    )
    notes_shape = _governance_notes_shape(slide, title_shape)
    if notes_shape and text.strip():
        _set_text_frame(notes_shape.text_frame, text, notes_shape, max_bullets=5)
    elif notes_shape:
        notes_shape.text_frame.clear()


def _parse_plan_phases(text: str) -> list[tuple[str, int, int]]:
    """Parse project plan lines into (label, start_week, end_week) 1-based inclusive."""
    phases: list[tuple[str, int, int]] = []
    for line in _bullet_lines(text):
        if "|" not in line:
            continue
        label, span = line.split("|", 1)
        label = label.strip()
        span = span.strip().lower()
        start, end = 1, 2
        if "week" in span:
            digits = [int(part) for part in re.findall(r"\d+", span)]
            if len(digits) >= 2:
                start, end = digits[0], digits[1]
            elif len(digits) == 1:
                start, end = digits[0], digits[0]
        else:
            marks = [index + 1 for index, cell in enumerate(span.split("|")) if cell.strip()]
            if marks:
                start, end = marks[0], marks[-1]
        phases.append((label, start, end))
    return phases


def _default_project_plan_phases(proposal: dict[str, Any]) -> list[tuple[str, int, int]]:
    impl_lines = _bullet_lines(_coerce_text(proposal.get("implementationMethodology", "")))
    labels = [
        impl_lines[0] if impl_lines else "Discovery & Design",
        impl_lines[1] if len(impl_lines) > 1 else "Iterative Development",
        impl_lines[2] if len(impl_lines) > 2 else "UAT & Go-Live",
        impl_lines[3] if len(impl_lines) > 3 else "Hypercare Support",
    ]
    spans = [(1, 2), (3, 9), (10, 12), (13, 15)]
    return list(zip(labels, [s[0] for s in spans], [s[1] for s in spans]))


def _update_project_plan_slide(slide: Any, proposal: dict[str, Any]) -> None:
    _ensure_slide_title(slide, SLIDE_TITLES[PROJECT_PLAN_SLIDE_INDEX])
    tables = _table_shapes(slide)
    if not tables:
        return

    table = tables[0].table
    plan_text = _coerce_text(proposal.get("projectPlan", ""))
    phases = _parse_plan_phases(plan_text)
    if not phases:
        phases = _default_project_plan_phases(proposal)

    for row_index in range(1, len(table.rows)):
        phase_index = row_index - 1
        if phase_index < len(phases):
            label, start_week, end_week = phases[phase_index]
            table.cell(row_index, 0).text = _truncate_text(label, 45)
            for week_col in range(1, len(table.columns)):
                week_num = week_col
                mark = "●" if start_week <= week_num <= end_week else ""
                table.cell(row_index, week_col).text = mark
        else:
            for col_index in range(len(table.columns)):
                table.cell(row_index, col_index).text = ""

    for shape in _iter_text_shapes(slide):
        if getattr(shape, "top", 0) > 4_200_000 and "Flutter" in shape.text:
            shape.text = (
                "Timeline aligned to agreed scope; exact durations confirmed during discovery."
            )
            _apply_body_font(shape.text_frame)


def _update_three_column_table(table: Any, text: str) -> bool:
    lines = _content_to_lines(text)
    filled = False
    for row_index in range(1, len(table.rows)):
        line_index = row_index - 1
        if line_index < len(lines):
            parts = [part.strip() for part in lines[line_index].split("|")]
            if len(parts) >= 3:
                for col_index in range(min(3, len(table.columns))):
                    table.cell(row_index, col_index).text = _truncate_text(
                        parts[col_index], 100
                    )
                filled = True
            else:
                for col_index in range(len(table.columns)):
                    table.cell(row_index, col_index).text = ""
        else:
            for col_index in range(len(table.columns)):
                table.cell(row_index, col_index).text = ""
    return filled


def _update_slide_table(slide: Any, text: str) -> bool:
    tables = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE]
    if not tables:
        return False

    table = tables[0].table
    if len(table.columns) >= 3:
        if not _update_three_column_table(table, text):
            _clear_table_data_rows(table)
            return False
        return True

    return _update_two_column_table(table, text)


def _clear_center_overlay_text(slide: Any, title_shape: Any | None) -> None:
    """Remove paragraph dumps that overlap diagram layouts."""
    for shape in _iter_text_shapes(slide):
        if shape is title_shape or _is_footer_shape(shape):
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            continue
        width = getattr(shape, "width", 0)
        text = shape.text.strip()
        if width > 3_500_000 and len(text) > 35:
            shape.text_frame.clear()


def _find_diagram_annotation_shape(slide: Any, title_shape: Any | None) -> Any | None:
    candidates: list[tuple[int, Any]] = []
    for shape in _iter_text_shapes(slide):
        if shape is title_shape or _is_footer_shape(shape):
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            continue
        capacity = _shape_char_capacity(shape)
        if capacity < 60 or capacity > 500:
            continue
        if getattr(shape, "left", 0) < 3_500_000:
            continue
        candidates.append((capacity, shape))
    if not candidates:
        return None
    return max(candidates, key=lambda item: item[0])[1]


def _update_diagram_slide(slide: Any, text: str, title: str) -> None:
    _ensure_slide_title(slide, title)
    shapes = list(_iter_text_shapes(slide))
    title_shape = _find_title_shape(shapes)
    _clear_center_overlay_text(slide, title_shape)

    annotation_shape = _find_diagram_annotation_shape(slide, title_shape)
    if annotation_shape and text.strip():
        _set_text_frame(annotation_shape.text_frame, text, annotation_shape, max_bullets=5)
    elif annotation_shape and not text.strip():
        annotation_shape.text_frame.clear()


def _is_multi_card_layout(body_shapes: list[Any]) -> bool:
    cards = [shape for shape in body_shapes if _shape_char_capacity(shape) >= 80]
    if len(cards) < 3:
        return False

    areas = [_shape_area(shape) for shape in cards]
    average = sum(areas) / len(areas)
    if average <= 0:
        return False

    return all(abs(area - average) / average < 0.3 for area in areas)


def _update_multi_card_slide(body_shapes: list[Any], text: str) -> None:
    cards = sorted(
        [shape for shape in body_shapes if _shape_char_capacity(shape) >= 80],
        key=lambda shape: (shape.top, shape.left),
    )
    lines = _content_to_lines(text)

    for index, shape in enumerate(cards):
        if index < len(lines):
            _set_text_frame(shape.text_frame, lines[index], shape, max_bullets=6)
        else:
            shape.text_frame.clear()


def _shape_has_sample_content(shape: Any) -> bool:
    text = shape.text if hasattr(shape, "text") else ""
    if any(marker in text for marker in _TEMPLATE_CLIENT_MARKERS):
        return True
    return any(phrase in text for phrase in _TEMPLATE_SAMPLE_PHRASES)


def _clear_slide_body_shapes(slide: Any, keep_title: bool = True) -> None:
    shapes = list(_iter_text_shapes(slide))
    title_shape = _find_title_shape(shapes) if keep_title else None
    for shape in shapes:
        if keep_title and _is_title_shape(shape, title_shape):
            continue
        if _is_footer_shape(shape):
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            continue
        if _shape_char_capacity(shape) < 60:
            continue
        shape.text_frame.clear()


def _purge_template_sample_text(slide: Any) -> None:
    shapes = list(_iter_text_shapes(slide))
    title_shape = _find_title_shape(shapes)
    for shape in shapes:
        if _is_title_shape(shape, title_shape) or _is_footer_shape(shape):
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            continue
        if _shape_has_sample_content(shape):
            shape.text_frame.clear()


def _update_slide_body(slide: Any, text: str, slide_index: int | None = None) -> None:
    if not text.strip():
        return

    _purge_template_sample_text(slide)

    text_shapes = list(_iter_text_shapes(slide))
    title_shape = _find_title_shape(text_shapes)
    body_shapes: list[Any] = []

    for shape in text_shapes:
        if _is_footer_shape(shape) or _is_title_shape(shape, title_shape):
            continue
        if _shape_char_capacity(shape) <= 0:
            continue
        body_shapes.append(shape)

    if not body_shapes:
        return

    force_cards = slide_index in MULTI_CARD_SLIDE_INDICES
    if force_cards or _is_multi_card_layout(body_shapes):
        cards = [shape for shape in body_shapes if _shape_char_capacity(shape) >= 80]
        for shape in body_shapes:
            if shape not in cards:
                shape.text_frame.clear()
        _update_multi_card_slide(body_shapes, text)
        cards = [shape for shape in body_shapes if _shape_char_capacity(shape) >= 80]
        if cards and not any(_shape_has_visible_text(shape) for shape in cards):
            primary_shape = max(body_shapes, key=_shape_area)
            _set_text_frame(primary_shape.text_frame, text, primary_shape)
        return

    primary_shape = max(body_shapes, key=_shape_area)
    for shape in body_shapes:
        if shape is not primary_shape:
            shape.text_frame.clear()

    _set_text_frame(primary_shape.text_frame, text, primary_shape)


def _update_slide_content(
    slide: Any,
    text: str,
    slide_index: int | None = None,
) -> None:
    title = SLIDE_TITLES.get(slide_index or -1, "")
    if title:
        _ensure_slide_title(slide, title)

    title_shape = _find_title_shape(list(_iter_text_shapes(slide)))

    if slide_index in DIAGRAM_OVERLAY_SLIDE_INDICES:
        _update_diagram_slide(slide, text, title)
        return

    if slide_index == TABLE_COMPARISON_SLIDE_INDEX:
        comparison = _coerce_text(text)
        if not _update_slide_table(slide, comparison):
            _clear_center_overlay_text(slide, title_shape)
        _ensure_slide_title(slide, SLIDE_TITLES[TABLE_COMPARISON_SLIDE_INDEX])
        return

    if slide_index in TABLE_TWO_COLUMN_SLIDES:
        if _update_slide_table(slide, text):
            _clear_footnotes(slide, title_shape)
            _ensure_slide_title(slide, title)
            return

    if _update_slide_table(slide, text):
        _clear_footnotes(slide, title_shape)
        _ensure_slide_title(slide, title)
        return

    _update_slide_body(slide, text, slide_index=slide_index)
    _ensure_slide_title(slide, title)


def _bullet_lines(text: str) -> list[str]:
    return [line.strip() for line in text.split("\n") if line.strip()]


def _compliance_summary_text(proposal: dict[str, Any], max_items: int = 6) -> str:
    matrix = proposal.get("complianceMatrix") or []
    lines: list[str] = []
    for row in matrix[:max_items]:
        if not isinstance(row, dict):
            continue
        req_id = _coerce_text(row.get("requirementId", ""))
        response = _truncate_text(_coerce_text(row.get("response", "")), 140)
        if response:
            prefix = f"{req_id}: " if req_id else ""
            lines.append(f"• {prefix}{response}")
    return "\n".join(lines)


def _slide_has_meaningful_body(slide: Any) -> bool:
    shapes = list(_iter_text_shapes(slide))
    title_shape = _find_title_shape(shapes)
    for shape in shapes:
        if _is_footer_shape(shape) or _is_title_shape(shape, title_shape):
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            table = shape.table
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip() and cell.text.strip() not in SKIP_TEXT:
                        return True
            continue
        text = shape.text.strip() if hasattr(shape, "text") else ""
        if text and text not in SKIP_TEXT and len(text) > 24:
            return True
    return False


def _minimal_placeholder(slide_index: int, proposal: dict[str, Any]) -> str:
    customer = _customer_name(proposal)
    title = SLIDE_TITLES.get(slide_index, "this engagement")
    return (
        f"TTN will tailor {title.lower()} for {customer} based on the analysed "
        "requirements, agreed scope, and discovery outcomes."
    )


def _resolve_slide_text(proposal: dict[str, Any], slide_index: int) -> str:
    keys: list[str] = []
    primary = SLIDE_CONTENT_MAP.get(slide_index)
    if primary:
        keys.append(primary)
    keys.extend(SLIDE_CONTENT_FALLBACKS.get(slide_index, ()))

    seen: set[str] = set()
    for key in keys:
        if key in seen:
            continue
        seen.add(key)
        text = _coerce_text(proposal.get(key, ""))
        if text.strip():
            return text

    if slide_index in {2, 25}:
        summary = _compliance_summary_text(proposal)
        if summary:
            return summary
    return ""


def _shape_has_visible_text(shape: Any) -> bool:
    return bool(shape.text.strip()) if hasattr(shape, "text") else False


def _enrich_proposal_defaults(proposal: dict[str, Any]) -> None:
    """Avoid blank key slides when the LLM omits a section."""
    understanding = _coerce_text(proposal.get("understandingOfRequirements", ""))
    understanding_lines = _bullet_lines(understanding)

    summary = _coerce_text(proposal.get("executiveSummary", ""))
    if not summary:
        if understanding_lines:
            proposal["executiveSummary"] = "\n".join(understanding_lines[:4])
        elif understanding:
            proposal["executiveSummary"] = _truncate_text(understanding, 900)
        else:
            compliance_summary = _compliance_summary_text(proposal, max_items=4)
            if compliance_summary:
                proposal["executiveSummary"] = compliance_summary

    solution = _coerce_text(proposal.get("proposedSolution", ""))
    if not solution:
        if len(understanding_lines) > 4:
            proposal["proposedSolution"] = "\n".join(understanding_lines[4:12])
        elif understanding:
            proposal["proposedSolution"] = understanding

    tech = _coerce_text(proposal.get("technicalApproach", ""))
    if not tech:
        source = _coerce_text(proposal.get("proposedSolution", "")) or understanding
        proposal["technicalApproach"] = _truncate_text(source, 1200)

    if not _coerce_text(proposal.get("securityCompliance", "")):
        proposal["securityCompliance"] = _truncate_text(
            _coerce_text(proposal.get("supportAndSla", ""))
            or _coerce_text(proposal.get("technicalApproach", "")),
            1000,
        ) or (
            "Security and compliance aligned to organisational policies, "
            "data protection requirements, and industry standards."
        )

    if not _coerce_text(proposal.get("assumptions", "")):
        proposal["assumptions"] = (
            "Client provides timely access to stakeholders, systems, and environments.\n"
            "Requirements are baselined after discovery; material changes follow change control.\n"
            "Third-party dependencies and licenses are provided or procured as agreed."
        )

    impl = _coerce_text(proposal.get("implementationMethodology", ""))
    if not impl:
        proposal["implementationMethodology"] = _truncate_text(
            proposal.get("technicalApproach", ""), 800
        )

    support = _coerce_text(proposal.get("supportAndSla", ""))
    if not support:
        proposal["supportAndSla"] = (
            "Dedicated project governance with defined SLAs, escalation paths, "
            "and ongoing support aligned to agreed service levels."
        )

    if not _coerce_text(proposal.get("governanceModel", "")):
        proposal["governanceModel"] = _truncate_text(support, 600)

    if not _coerce_text(proposal.get("architectureOverview", "")):
        proposal["architectureOverview"] = _truncate_text(
            proposal.get("technicalApproach", ""), 1000
        )

    if not _coerce_text(proposal.get("mobilePlatformComparison", "")):
        proposal["mobilePlatformComparison"] = _comparison_lines(proposal)

    if not _coerce_text(proposal.get("projectPlan", "")):
        phases = _default_project_plan_phases(proposal)
        proposal["projectPlan"] = "\n".join(
            f"{label} | weeks {start}-{end}" for label, start, end in phases
        )


def _project_title(proposal: dict[str, Any]) -> str:
    return _coerce_text(
        proposal.get("projectTitle")
        or proposal.get("project_title")
        or proposal.get("customer")
        or "Client"
    )


def _customer_name(proposal: dict[str, Any]) -> str:
    return _coerce_text(proposal.get("customer") or _project_title(proposal))


def _slide_plain_text(slide: Any) -> str:
    return "\n".join(
        shape.text
        for shape in _iter_text_shapes(slide)
        if hasattr(shape, "text") and shape.text.strip()
    )


def _payment_milestone_slide_indices(prs: Presentation) -> list[int]:
    indices: list[int] = []
    for index, slide in enumerate(prs.slides):
        text = _slide_plain_text(slide)
        if "Payment Milestones" in text or (
            "Milestone" in text and "AUD" in text and "Price" in text
        ):
            indices.append(index)
    return sorted(indices, reverse=True)


def _apply_customer_substitutions(slide: Any, proposal: dict[str, Any]) -> None:
    customer = _customer_name(proposal)
    project = _project_title(proposal)
    replacements = [
        ("Racing Queensland Networks", customer),
        ("Racing Queensland", customer),
        ("Mobile App Development | Racing Queensland v1.0", f"TTN Proposal | {project}"),
        ("TTN Proposal for Mobile App Development | Racing Queensland v1.0", f"TTN Proposal | {project}"),
    ]
    for shape in _iter_text_shapes(slide):
        if _is_footer_shape(shape):
            continue
        text = shape.text
        for old, new in replacements:
            if old in text and new:
                text = text.replace(old, new)
        for marker in _TEMPLATE_CLIENT_MARKERS:
            if marker in text and customer:
                text = text.replace(marker, customer)
        if text != shape.text:
            shape.text = text


def _clear_template_bodies(slide: Any) -> None:
    _purge_template_sample_text(slide)
    _clear_slide_body_shapes(slide, keep_title=True)


def _backfill_sparse_content_slides(
    prs: Presentation, proposal: dict[str, Any]
) -> None:
    """Fill title-only slides after the main pass (e.g. when template cards were cleared)."""
    indices = (
        set(SLIDE_CONTENT_MAP)
        | set(TABLE_TWO_COLUMN_SLIDES)
        | {IN_SCOPE_SLIDE_INDEX, ENGAGEMENT_SLIDE_INDEX}
    )
    for index in sorted(indices):
        if index in FIXED_SLIDE_INDICES or index >= len(prs.slides):
            continue
        slide = prs.slides[index]
        if _slide_has_meaningful_body(slide):
            continue
        text = _resolve_slide_text(proposal, index)
        if not text:
            text = _minimal_placeholder(index, proposal)
        _update_slide_content(slide, text, slide_index=index)


def _update_agenda_slide(slide: Any, proposal: dict[str, Any]) -> None:
    customer = _customer_name(proposal)
    agenda_lines = [
        f"Executive Summary and Our Understanding of {customer} Requirements",
        "Proposed Technical Solution Approach",
        "Non-Functional Requirements and Architecture",
        "Implementation, Governance, and Project Plan",
        "Engagement Model and Commercial Summary",
        "Success Stories",
    ]
    body_shapes = [
        shape
        for shape in sorted(_iter_text_shapes(slide), key=lambda item: item.top)
        if _shape_char_capacity(shape) >= 40
        and shape.text.strip() not in SKIP_TEXT
        and shape.text.strip() != "Agenda"
    ]
    for index, shape in enumerate(body_shapes):
        if index >= len(agenda_lines):
            shape.text_frame.clear()
            continue
        shape.text = agenda_lines[index]
        _apply_body_font(shape.text_frame)


def _scope_column_shapes(slide: Any) -> list[Any]:
    title_shape = _find_title_shape(list(_iter_text_shapes(slide)))
    columns = [
        shape
        for shape in _iter_text_shapes(slide)
        if shape is not title_shape
        and not _is_footer_shape(shape)
        and getattr(shape, "top", 0) > 1_200_000
        and _shape_char_capacity(shape) >= 300
    ]
    return sorted(columns, key=lambda shape: shape.left)[:2]


def _update_scope_slide(slide: Any, proposal: dict[str, Any]) -> None:
    in_scope = _coerce_text(proposal.get("inScope", ""))
    out_of_scope = _coerce_text(proposal.get("outOfScope", ""))
    columns = _scope_column_shapes(slide)
    if len(columns) >= 2:
        if in_scope:
            _set_text_frame(columns[0].text_frame, in_scope, columns[0])
        else:
            columns[0].text_frame.clear()
        if out_of_scope:
            _set_text_frame(columns[1].text_frame, out_of_scope, columns[1])
        else:
            columns[1].text_frame.clear()

    title_shape = _find_title_shape(list(_iter_text_shapes(slide)))
    for shape in _iter_text_shapes(slide):
        if shape in columns or shape is title_shape or _is_footer_shape(shape):
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            continue
        text = shape.text.strip()
        if not text or text in {"In Scope", "Not Estimated", "In Scope and Out of Scope"}:
            continue
        shape.text_frame.clear()


def _update_engagement_slide(slide: Any, proposal: dict[str, Any]) -> None:
    model = _coerce_text(proposal.get("engagementModel", ""))
    commercial = _coerce_text(proposal.get("commercialSummary", "")) or (
        "Commercial pricing and payment terms will be agreed based on "
        "final scope, milestones, and contract."
    )
    columns = sorted(
        [
            shape
            for shape in _iter_text_shapes(slide)
            if _shape_char_capacity(shape) >= 400
        ],
        key=lambda shape: shape.left,
    )
    if columns:
        if model:
            _set_text_frame(columns[0].text_frame, model, columns[0])
        else:
            columns[0].text_frame.clear()
    if len(columns) >= 2:
        _set_text_frame(columns[1].text_frame, commercial, columns[1])


def _update_cover_slide(slide: Any, proposal: dict[str, Any], rfp_id: str) -> None:
    project = _project_title(proposal)
    date_str = datetime.now().strftime("%d %b %Y")
    title_line = f"TTN Proposal | {project}"

    for shape in slide.shapes:
        if not hasattr(shape, "text"):
            continue
        text = shape.text
        if any(
            token in text
            for token in ("Proposal", "Date", "Racing", "Mobile App Development")
        ):
            shape.text = f"{title_line}\nDate : {date_str}\nRef : {rfp_id}"


def render_pptx_from_template(
    proposal: dict[str, Any],
    output_path: str,
    rfp_id: str,
) -> str:
    template_path = resolve_pptx_template_path()
    if not template_path.exists():
        raise FileNotFoundError(
            f"PPTX template not found at {template_path}. "
            "Place proposal-template.pptx in templates/."
        )

    os_dir = Path(output_path).parent
    os_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)

    prs = Presentation(output_path)
    _enrich_proposal_defaults(proposal)

    for index, slide in enumerate(prs.slides):
        if index in FIXED_SLIDE_INDICES:
            continue

        if index == 0:
            _update_cover_slide(slide, proposal, rfp_id)
            continue

        if index in AGENDA_SLIDE_INDICES:
            _update_agenda_slide(slide, proposal)
            _apply_customer_substitutions(slide, proposal)
            continue

        if index == IN_SCOPE_SLIDE_INDEX:
            _ensure_slide_title(slide, SLIDE_TITLES[IN_SCOPE_SLIDE_INDEX])
            _update_scope_slide(slide, proposal)
            _apply_customer_substitutions(slide, proposal)
            continue

        if index == ENGAGEMENT_SLIDE_INDEX:
            _ensure_slide_title(slide, SLIDE_TITLES[ENGAGEMENT_SLIDE_INDEX])
            _update_engagement_slide(slide, proposal)
            _apply_customer_substitutions(slide, proposal)
            continue

        if index in DIAGRAM_ANNOTATION_CONTENT:
            key = DIAGRAM_ANNOTATION_CONTENT[index]
            content = _coerce_text(proposal.get(key, ""))
            _update_diagram_slide(
                slide,
                content,
                SLIDE_TITLES.get(index, ""),
            )
            _apply_customer_substitutions(slide, proposal)
            continue

        if index in TABLE_TWO_COLUMN_SLIDES:
            content = _resolve_slide_text(proposal, index)
            _update_slide_content(slide, content, slide_index=index)
            if not _slide_has_meaningful_body(slide) and content:
                _update_slide_body(slide, content, slide_index=index)
            _apply_customer_substitutions(slide, proposal)
            continue

        if index == TABLE_COMPARISON_SLIDE_INDEX:
            _update_comparison_slide(slide, proposal)
            _apply_customer_substitutions(slide, proposal)
            continue

        if index == ARCHITECTURE_SLIDE_INDEX:
            _update_architecture_slide(slide, proposal)
            _apply_customer_substitutions(slide, proposal)
            continue

        if index == GOVERNANCE_SLIDE_INDEX:
            _update_governance_slide(slide, proposal)
            continue

        if index == PROJECT_PLAN_SLIDE_INDEX:
            _update_project_plan_slide(slide, proposal)
            _apply_customer_substitutions(slide, proposal)
            continue

        if index in DIAGRAM_OVERLAY_SLIDE_INDICES:
            _update_diagram_slide(slide, "", SLIDE_TITLES.get(index, ""))
            _apply_customer_substitutions(slide, proposal)
            continue

        content_key = SLIDE_CONTENT_MAP.get(index)
        content = _resolve_slide_text(proposal, index) if content_key else ""
        if content_key:
            _purge_template_sample_text(slide)
        if content:
            _update_slide_content(slide, content, slide_index=index)
        elif content_key:
            _ensure_slide_title(slide, SLIDE_TITLES.get(index, ""))
            placeholder = _minimal_placeholder(index, proposal)
            _update_slide_content(slide, placeholder, slide_index=index)

        _apply_customer_substitutions(slide, proposal)

    _backfill_sparse_content_slides(prs, proposal)

    if (
        len(prs.slides) > EMPTY_TRAILING_SLIDE_INDEX
        and not _slide_has_meaningful_body(prs.slides[EMPTY_TRAILING_SLIDE_INDEX])
    ):
        _delete_slide(prs, EMPTY_TRAILING_SLIDE_INDEX)

    for index in _payment_milestone_slide_indices(prs):
        _delete_slide(prs, index)

    prs.save(output_path)
    return output_path


def _delete_slide(prs: Presentation, index: int) -> None:
    """Remove a slide by 0-based index (python-pptx has no public delete API)."""
    slide_id_list = prs.slides._sldIdLst
    slide_ids = list(slide_id_list)
    if index < 0 or index >= len(slide_ids):
        return
    r_id = slide_ids[index].rId
    prs.part.drop_rel(r_id)
    slide_id_list.remove(slide_ids[index])
