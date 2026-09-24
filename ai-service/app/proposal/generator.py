import asyncio
import os
import shutil
import subprocess
from functools import partial
from pathlib import Path
from typing import Any, Callable

from docx import Document

from app.config import settings
from app.proposal.pptx_template import render_pptx_from_template
from app.providers.ollama_provider import ollama_provider

PROPOSAL_SYSTEM = """You are a proposal writer. Write concise, professional proposal sections from extracted RFP requirements and reviewed responses.
Ground every section in the supplied requirements (from uploaded documents and/or crawled website pages) and the response evidence.
Use only the provided material. Do not invent capabilities. Return valid JSON only.
Format for PowerPoint: use newline-separated bullets; keep lines under 120 characters where possible.
For technicalApproach, include some lines as "Component | Key consideration" for architecture tables.
For mobilePlatformComparison: base the comparison on the actual requirements.
- If requirements are website/web-portal and no technology stack is stated, use columns "Next.js (Frontend)" and "Drupal (Backend)" (never Flutter/React Native).
- Only use Flutter vs React Native when requirements clearly call for a native/mobile app.
- Format 5-8 lines as "Factor | Option A | Option B" (three pipe-separated parts).
For architectureOverview use short layer labels (one line each) for UI, APIs, data, integrations.
For governanceModel use 4-6 short bullets on cadence, steering, and escalation (no long paragraphs).
For projectPlan use 4 lines: "Phase name | weeks 1-3" (inclusive week ranges on a 15-week plan)."""

PROPOSAL_BATCHES: list[list[str]] = [
    ["executiveSummary", "understandingOfRequirements", "proposedSolution"],
    [
        "technicalApproach",
        "architectureOverview",
        "mobilePlatformComparison",
        "implementationMethodology",
        "projectPlan",
        "governanceModel",
        "supportAndSla",
    ],
    [
        "securityCompliance",
        "assumptions",
        "inScope",
        "outOfScope",
        "engagementModel",
    ],
]

SECTION_LABELS: dict[str, str] = {
    "executiveSummary": "Executive Summary",
    "understandingOfRequirements": "Understanding of Requirements",
    "proposedSolution": "Proposed Solution",
    "technicalApproach": "Technical Approach",
    "architectureOverview": "Architecture Overview",
    "mobilePlatformComparison": "Technology Platform Comparison",
    "implementationMethodology": "Implementation Methodology",
    "projectPlan": "Project Plan",
    "governanceModel": "Governance Model",
    "supportAndSla": "Support & SLA",
    "securityCompliance": "Security & Compliance",
    "assumptions": "Assumptions",
    "inScope": "In Scope",
    "outOfScope": "Out of Scope",
    "engagementModel": "Engagement Model",
}

ProgressFn = Callable[..., None]


def _coerce_text(value: Any, max_len: int | None = None) -> str:
    """Normalize LLM output (string, list, or nested dict) into plain text."""
    if value is None:
        text = ""
    elif isinstance(value, str):
        text = value
    elif isinstance(value, list):
        parts = [_coerce_text(item) for item in value]
        text = "\n".join(part for part in parts if part)
    elif isinstance(value, dict):
        for key in ("text", "content", "body", "paragraph", "summary", "value"):
            if key in value:
                text = _coerce_text(value[key])
                break
        else:
            parts = [_coerce_text(item) for item in value.values()]
            text = "\n".join(part for part in parts if part)
    else:
        text = str(value)

    if max_len is not None:
        return text[:max_len]
    return text


def _proposal_model() -> str:
    return settings.proposal_model or settings.llm_fast_model


def _build_summary(responses: list[dict[str, Any]]) -> str:
    max_items = settings.proposal_max_responses
    max_chars = settings.proposal_response_chars
    lines: list[str] = []
    for resp in responses[:max_items]:
        req_id = resp.get("requirementId", "")
        text = (resp.get("response") or "")[:max_chars]
        status = resp.get("status", "")
        if text:
            lines.append(f"- {req_id} ({status}): {text}")
    return "\n".join(lines) or "- No reviewed responses available."


def _build_requirements_digest(requirements: list[dict[str, Any]]) -> str:
    max_items = settings.proposal_max_responses
    max_chars = settings.proposal_response_chars
    lines: list[str] = []
    for req in requirements[:max_items]:
        req_id = req.get("id") or req.get("req_id") or ""
        description = (req.get("description") or "")[:max_chars]
        if not description:
            continue
        mandatory = req.get("mandatory", False)
        req_type = req.get("type") or ""
        source_type = req.get("source_type") or ""
        source_section = (req.get("source_section") or "").strip()
        if not source_type and source_section:
            source_type = (
                "website"
                if source_section.startswith("/") or "website" in source_section.lower()
                else "document"
            )
        prefix = "[MANDATORY] " if mandatory else ""
        source_hint = f" ({source_type})" if source_type else ""
        lines.append(f"- {req_id}{source_hint} [{req_type}] {prefix}{description}")
    return "\n".join(lines) or "- No requirements were extracted."


def _format_source_context(metadata: dict[str, Any]) -> str:
    sources = metadata.get("sources")
    if not isinstance(sources, dict):
        return "RFP input: uploaded document and/or website URL."

    parts: list[str] = []
    document = sources.get("document")
    if isinstance(document, dict) and document.get("filename"):
        parts.append(f"Uploaded document: {document['filename']}")
    website = sources.get("website")
    if isinstance(website, dict) and website.get("url"):
        pages = website.get("pages_crawled", "?")
        parts.append(f"Website crawl: {website['url']} ({pages} pages)")
    return "\n".join(parts) or "RFP input: uploaded document and/or website URL."


def _progress(
    on_progress: ProgressFn | None,
    percent: int,
    step: str,
    message: str,
    **extra: Any,
) -> None:
    if on_progress:
        on_progress(percent=percent, step=step, message=message, **extra)


async def _generate_batch(
    keys: list[str],
    summary: str,
    requirements_digest: str,
    metadata: dict[str, Any],
) -> dict[str, str]:
    key_list = ", ".join(f'"{k}"' for k in keys)
    content = await ollama_provider.chat(
        messages=[
            {"role": "system", "content": PROPOSAL_SYSTEM},
            {
                "role": "user",
                "content": f"""Write these proposal sections from the extracted requirements and reviewed responses.

Customer: {metadata.get("customer", "Client")}
Industry: {metadata.get("industry", "General")}
Input sources:
{_format_source_context(metadata)}

Extracted requirements (document upload and/or website crawl):
{requirements_digest}

Reviewed responses:
{summary}

Return JSON with ONLY these keys: {key_list}
Each value MUST be a plain string (bullet lines separated by newlines). Do not nest objects or arrays.

Example shape: {{{", ".join(f'"{k}": "bullet one\\nbullet two"' for k in keys)}}}""",
            },
        ],
        model=_proposal_model(),
        format_json=True,
        disable_thinking=True,
        max_tokens=settings.proposal_max_tokens,
    )

    try:
        parsed = ollama_provider.parse_json_response(content)
        if isinstance(parsed, dict):
            return {k: _coerce_text(parsed.get(k, "")) for k in keys}
    except (ValueError, TypeError):
        pass

    return {k: "" for k in keys}


async def generate_proposal_content(
    metadata: dict[str, Any],
    requirements: list[dict[str, Any]],
    responses: list[dict[str, Any]],
    compliance: dict[str, Any],
    on_progress: ProgressFn | None = None,
) -> dict[str, Any]:
    summary = _build_summary(responses)
    requirements_digest = _build_requirements_digest(requirements)
    total_batches = len(PROPOSAL_BATCHES)
    done_batches = 0
    proposal: dict[str, Any] = {}
    sem = asyncio.Semaphore(max(1, settings.proposal_concurrency))

    async def run_batch(batch_idx: int, keys: list[str]) -> dict[str, str]:
        nonlocal done_batches
        labels = [SECTION_LABELS.get(k, k) for k in keys]
        async with sem:
            _progress(
                on_progress,
                5 + int((done_batches / total_batches) * 70),
                "writing",
                f"Writing: {', '.join(labels[:2])}{'…' if len(labels) > 2 else ''}",
                sectionDone=done_batches,
                sectionTotal=total_batches,
            )
            result = await _generate_batch(keys, summary, requirements_digest, metadata)
            done_batches += 1
            _progress(
                on_progress,
                5 + int((done_batches / total_batches) * 70),
                "writing",
                f"Completed {done_batches}/{total_batches} section groups",
                sectionDone=done_batches,
                sectionTotal=total_batches,
            )
            return result

    batch_results = await asyncio.gather(
        *[run_batch(i, keys) for i, keys in enumerate(PROPOSAL_BATCHES)]
    )
    for batch in batch_results:
        proposal.update(batch)

    _progress(on_progress, 80, "render", "Building Word document…")
    proposal["complianceMatrix"] = compliance.get("complianceMatrix", [])
    proposal["customer"] = metadata.get("customer", "")
    proposal["projectTitle"] = (
        metadata.get("projectTitle")
        or metadata.get("project_title")
        or metadata.get("customer")
        or "Client"
    )
    if isinstance(metadata.get("sources"), dict):
        proposal["sources"] = metadata["sources"]
    # Help platform-comparison heuristics when the LLM omits stack cues.
    if requirements_digest and not _coerce_text(proposal.get("understandingOfRequirements")):
        proposal["understandingOfRequirements"] = requirements_digest[:1500]
    elif requirements_digest:
        # Append digest fragment so website/mobile detection sees raw requirements.
        proposal["_requirementsDigest"] = requirements_digest[:2000]
    return proposal


def render_docx(proposal: dict[str, Any], output_path: str, rfp_id: str) -> str:
    """Render structured proposal JSON into a DOCX file."""
    template = Path(settings.template_path)
    if template.exists():
        doc = Document(str(template))
    else:
        doc = Document()

    doc.add_heading("RFP Response Proposal", 0)
    doc.add_paragraph(f"RFP Reference: {rfp_id}")
    doc.add_paragraph(f"Customer: {_coerce_text(proposal.get('customer', 'N/A'))}")

    sections = [
        ("Executive Summary", "executiveSummary"),
        ("Understanding of Requirements", "understandingOfRequirements"),
        ("Proposed Solution", "proposedSolution"),
        ("Technical Approach", "technicalApproach"),
        ("Implementation Methodology", "implementationMethodology"),
        ("Support and SLA", "supportAndSla"),
        ("Security and Compliance", "securityCompliance"),
        ("Assumptions", "assumptions"),
        ("In Scope", "inScope"),
        ("Out of Scope", "outOfScope"),
        ("Engagement Model", "engagementModel"),
    ]

    for title, key in sections:
        text = _coerce_text(proposal.get(key, ""))
        if text:
            doc.add_heading(title, level=1)
            doc.add_paragraph(text)

    matrix = proposal.get("complianceMatrix", [])
    if matrix:
        doc.add_heading("Compliance Matrix", level=1)
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        headers = ["Req ID", "Requirement", "Response", "Status", "Review"]
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = h

        for row in matrix[:100]:
            cells = table.add_row().cells
            cells[0].text = _coerce_text(row.get("requirementId", ""))
            cells[1].text = _coerce_text(row.get("requirement"), max_len=200)
            cells[2].text = _coerce_text(row.get("response"), max_len=300)
            cells[3].text = _coerce_text(row.get("status", ""))
            cells[4].text = "Yes" if row.get("reviewRequired") else "No"

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def render_pptx(proposal: dict[str, Any], output_path: str, rfp_id: str) -> str:
    """Render proposal using the TTN PPTX template."""
    return render_pptx_from_template(proposal, output_path, rfp_id)


def find_libreoffice_command() -> str | None:
    for command in ("libreoffice", "soffice"):
        resolved = shutil.which(command)
        if resolved:
            return resolved

    mac_path = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if os.path.isfile(mac_path):
        return mac_path

    return None


def convert_to_pdf(docx_path: str, force: bool = False) -> str | None:
    """Convert DOCX to PDF using LibreOffice headless."""
    if not force and not settings.proposal_generate_pdf:
        return None

    libreoffice = find_libreoffice_command()
    if not libreoffice:
        return None

    pdf_path = docx_path.replace(".docx", ".pdf")
    try:
        subprocess.run(
            [
                libreoffice,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                os.path.dirname(docx_path),
                docx_path,
            ],
            check=True,
            capture_output=True,
            timeout=settings.proposal_pdf_timeout_seconds,
        )
        return pdf_path if os.path.exists(pdf_path) else None
    except (subprocess.CalledProcessError, FileNotFoundError, subprocess.TimeoutExpired):
        return None


async def build_proposal_files(
    rfp_id: str,
    metadata: dict[str, Any],
    requirements: list[dict[str, Any]],
    responses: list[dict[str, Any]],
    compliance: dict[str, Any],
    on_progress: ProgressFn | None = None,
) -> dict[str, Any]:
    proposal = await generate_proposal_content(
        metadata, requirements, responses, compliance, on_progress=on_progress
    )

    docx_path = os.path.join(settings.proposal_dir, f"{rfp_id}_proposal.docx")
    loop = asyncio.get_running_loop()
    await loop.run_in_executor(
        None, partial(render_docx, proposal, docx_path, rfp_id)
    )

    _progress(on_progress, 88, "pptx", "Building PowerPoint deck…")
    pptx_path = os.path.join(settings.proposal_dir, f"{rfp_id}_proposal.pptx")
    await loop.run_in_executor(
        None, partial(render_pptx, proposal, pptx_path, rfp_id)
    )

    _progress(on_progress, 94, "pdf", "Converting to PDF (if available)…")
    pdf_path = await loop.run_in_executor(None, partial(convert_to_pdf, docx_path))

    _progress(on_progress, 100, "complete", "Proposal ready for download")
    return {
        "proposal": proposal,
        "docxPath": docx_path,
        "pptxPath": pptx_path,
        "pdfPath": pdf_path,
    }
