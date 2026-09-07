from pathlib import Path
from typing import Any

import fitz  # PyMuPDF
from docx import Document


def parse_pdf(file_path: str) -> dict[str, Any]:
    doc = fitz.open(file_path)
    pages: list[dict[str, Any]] = []
    full_text_parts: list[str] = []

    for page_num, page in enumerate(doc, start=1):
        text = page.get_text("text")
        pages.append({"page": page_num, "text": text})
        full_text_parts.append(text)

    metadata = doc.metadata or {}
    doc.close()

    return {
        "text": "\n\n".join(full_text_parts),
        "pages": pages,
        "metadata": {
            "title": metadata.get("title", ""),
            "author": metadata.get("author", ""),
            "page_count": len(pages),
        },
    }


def parse_docx(file_path: str) -> dict[str, Any]:
    doc = Document(file_path)
    paragraphs: list[dict[str, Any]] = []
    full_text_parts: list[str] = []

    for idx, para in enumerate(doc.paragraphs, start=1):
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name if para.style else "Normal"
        paragraphs.append({"index": idx, "text": text, "style": style})
        full_text_parts.append(text)

    tables: list[list[list[str]]] = []
    for table in doc.tables:
        rows = [
            [cell.text.strip() for cell in row.cells] for row in table.rows
        ]
        tables.append(rows)

    return {
        "text": "\n\n".join(full_text_parts),
        "paragraphs": paragraphs,
        "tables": tables,
        "metadata": {
            "title": doc.core_properties.title or "",
            "author": doc.core_properties.author or "",
            "paragraph_count": len(paragraphs),
        },
    }


def parse_document(file_path: str) -> dict[str, Any]:
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        result = parse_pdf(file_path)
        result["format"] = "pdf"
    elif suffix == ".docx":
        result = parse_docx(file_path)
        result["format"] = "docx"
    else:
        raise ValueError(f"Unsupported file format: {suffix}")

    result["filename"] = path.name
    return result
