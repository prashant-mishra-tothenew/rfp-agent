from docx import Document

from app.proposal.rfp_generator import render_rfp_docx


def test_render_rfp_docx(tmp_path):
    output_path = tmp_path / "generated-rfp.docx"
    rfp = {
        "title": "Cloud Platform RFP",
        "organization": "Example Corp",
        "industry": "Technology",
        "background": "Example Corp needs a managed cloud platform.",
        "objectives": ["Improve reliability"],
        "scope": ["Platform implementation"],
        "requirements": [
            {
                "id": "REQ-001",
                "category": "Technical",
                "description": "The platform must support Kubernetes.",
                "priority": "Mandatory",
            }
        ],
        "deliverables": ["Implementation plan"],
    }

    result = render_rfp_docx(rfp, str(output_path), "RFP-001")

    assert result == str(output_path)
    assert output_path.exists()

    document = Document(output_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Cloud Platform RFP" in text
    assert "Example Corp" in text
    assert document.tables[0].cell(1, 0).text == "REQ-001"
