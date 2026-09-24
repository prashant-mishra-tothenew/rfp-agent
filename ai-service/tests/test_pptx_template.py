import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.proposal.pptx_template import (
    FIXED_SLIDE_INDICES,
    SUCCESS_STORY_SLIDE_INDICES,
    WHY_TTN_SLIDE_INDEX,
    _shape_char_capacity,
    _truncate_text,
    render_pptx_from_template,
    resolve_pptx_template_path,
)


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip():
            parts.append(shape.text)
    return "\n".join(parts)


def test_template_path_exists():
    assert resolve_pptx_template_path().exists()


def test_sparse_proposal_does_not_leave_title_only_slides():
    with tempfile.TemporaryDirectory() as tmp:
        output = str(Path(tmp) / "proposal.pptx")
        render_pptx_from_template(
            {
                "customer": "Invest India",
                "complianceMatrix": [
                    {
                        "requirementId": "REQ-001",
                        "response": "Public portal for investment opportunities.",
                        "status": "Partial",
                    },
                    {
                        "requirementId": "REQ-002",
                        "response": "Search across sectors and programmes.",
                        "status": "Not supported",
                    },
                ],
            },
            output,
            "rfp-sparse-001",
        )
        prs = Presentation(output)
        assert len(prs.slides) == 26
        for index in (2, 5, 7, 11, 25):
            text = _slide_text(prs.slides[index])
            assert len(text) > 40
            assert "‹#›" not in text or text.replace("‹#›", "").strip()


def test_render_preserves_why_ttn_and_success_slides():
    template_path = resolve_pptx_template_path()
    template_prs = Presentation(str(template_path))
    why_ttn_before = _slide_text(template_prs.slides[WHY_TTN_SLIDE_INDEX])
    success_index = min(SUCCESS_STORY_SLIDE_INDICES)
    success_before = _slide_text(template_prs.slides[success_index])

    with tempfile.TemporaryDirectory() as tmp:
        output = str(Path(tmp) / "proposal.pptx")
        render_pptx_from_template(
            {
                "customer": "Acme Corp",
                "projectTitle": "Acme Mobile Platform",
                "executiveSummary": "Custom executive summary for Acme.",
                "understandingOfRequirements": "Custom requirements understanding.",
                "proposedSolution": "Custom proposed solution for Acme retail platform.",
                "technicalApproach": "API-first architecture | Microservices on AWS EKS",
                "engagementModel": "Iterative agile delivery with fortnightly demos.",
                "inScope": "Mobile app delivery\nAPI integration",
                "outOfScope": "Hardware procurement\nThird-party license costs",
            },
            output,
            "rfp-test-001",
        )

        result_prs = Presentation(output)
        assert _slide_text(result_prs.slides[WHY_TTN_SLIDE_INDEX]) == why_ttn_before
        assert _slide_text(result_prs.slides[success_index]) == success_before
        cover = _slide_text(result_prs.slides[0])
        assert "Acme Mobile Platform" in cover
        assert "Racing Queensland" not in cover
        assert "Custom executive summary" in _slide_text(result_prs.slides[2])
        slide5_text = _slide_text(result_prs.slides[5])
        assert "Custom proposed solution" in slide5_text
        assert "Races & Profiles" not in slide5_text
        scope_text = _slide_text(result_prs.slides[9])
        assert "Mobile app delivery" in scope_text
        assert "Hardware procurement" in scope_text
        assert "Procurement & Maintenance of Infrastructure" not in scope_text
        engagement_text = _slide_text(result_prs.slides[20])
        assert "Iterative agile delivery" in engagement_text
        assert "AUD 78,727" not in engagement_text


def test_truncate_text_respects_limit():
    long_text = (
        "Our proposed solution combines a Next.js/React frontend with REST/GraphQL APIs "
        "for a decoupled architecture. We integrate headless CMS platforms like Sanity."
    )
    truncated = _truncate_text(long_text, 90)
    assert len(truncated) <= 90
    assert len(truncated) > 40


def test_render_distributes_solution_cards_without_template_overflow():
    template_path = resolve_pptx_template_path()
    long_solution = "\n".join(
        [
            "Drupal Commerce core platform with multi-store support.",
            "Solr-powered faceted search for 100,000+ SKUs.",
            "Headless Next.js storefront with SSR for SEO.",
            "SAP ERP integration for inventory and orders.",
            "Stripe and PayPal checkout with PCI-DSS flows.",
            "WCAG 2.1 AA accessible customer journeys.",
        ]
    )

    with tempfile.TemporaryDirectory() as tmp:
        output = str(Path(tmp) / "proposal.pptx")
        render_pptx_from_template(
            {
                "customer": "Acme Corp",
                "proposedSolution": long_solution,
            },
            output,
            "rfp-test-002",
        )

        slide = Presentation(output).slides[5]
        card_capacities = [
            _shape_char_capacity(shape)
            for shape in slide.shapes
            if hasattr(shape, "text_frame") and _shape_char_capacity(shape) >= 100
        ]
        assert len(card_capacities) >= 3

        for shape in slide.shapes:
            if not hasattr(shape, "text_frame"):
                continue
            capacity = _shape_char_capacity(shape)
            if capacity < 100:
                continue
            assert len(shape.text) <= capacity + 5
            assert "Races & Profiles" not in shape.text


def test_architecture_slide_does_not_overlay_diagram_text():
    with tempfile.TemporaryDirectory() as tmp:
        output = str(Path(tmp) / "proposal.pptx")
        render_pptx_from_template(
            {
                "customer": "Acme Corp",
                "technicalApproach": "API Gateway | Federated GraphQL layer\nCDN | Global edge caching",
            },
            output,
            "rfp-test-arch",
        )
        slide = Presentation(output).slides[12]
        overlay = [
            shape.text
            for shape in slide.shapes
            if hasattr(shape, "text_frame")
            and getattr(shape, "width", 0) > 3_500_000
            and len(shape.text.strip()) > 40
        ]
        assert overlay == []
        assert _slide_text(slide).startswith("Proposed High Level Architecture")


def test_comparison_and_project_plan_slides_receive_content():
    with tempfile.TemporaryDirectory() as tmp:
        output = str(Path(tmp) / "proposal.pptx")
        render_pptx_from_template(
            {
                "customer": "Acme Corp",
                "technicalApproach": "Next.js storefront | Headless CMS integration",
                "implementationMethodology": "Discovery workshops\nAgile delivery sprints\nUAT and launch",
                "sources": {"website": {"url": "https://example.com", "pages_crawled": 5}},
                "understandingOfRequirements": "Public website portal with content and search.",
            },
            output,
            "rfp-plan-test",
        )
        prs = Presentation(output)
        comparison_table = next(
            sh.table
            for sh in prs.slides[14].shapes
            if sh.shape_type == MSO_SHAPE_TYPE.TABLE
        )
        header = " ".join(
            comparison_table.cell(0, col).text for col in range(3)
        ).lower()
        assert "next" in header
        assert "drupal" in header
        assert "flutter" not in header
        assert "react native" not in header
        assert comparison_table.cell(1, 0).text.strip()
        assert comparison_table.cell(1, 1).text.strip()

        plan_table = next(
            sh.table
            for sh in prs.slides[18].shapes
            if sh.shape_type == MSO_SHAPE_TYPE.TABLE
        )
        assert plan_table.cell(1, 0).text.strip()
        assert any(
            plan_table.cell(1, col).text.strip() == "●"
            for col in range(1, len(plan_table.columns))
        )


def test_website_rfp_rejects_mobile_comparison_defaults():
    with tempfile.TemporaryDirectory() as tmp:
        output = str(Path(tmp) / "proposal.pptx")
        render_pptx_from_template(
            {
                "customer": "Invest India",
                "sources": {
                    "website": {"url": "https://www.investindia.gov.in", "pages_crawled": 12}
                },
                "understandingOfRequirements": "Website features for investment opportunities.",
                "mobilePlatformComparison": (
                    "Factor | REACT NATIVE | Flutter\n"
                    "Platform | Web-based | Native (iOS/Android)\n"
                    "Development | Single codebase | Separate codebases"
                ),
            },
            output,
            "rfp-web-compare",
        )
        slide = Presentation(output).slides[14]
        titles = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text_frame")
            and shape.top < 900_000
            and shape.text.strip()
        ]
        assert any("Web Technology" in title for title in titles)
        comparison_table = next(
            sh.table
            for sh in slide.shapes
            if sh.shape_type == MSO_SHAPE_TYPE.TABLE
        )
        header = " ".join(
            comparison_table.cell(0, col).text for col in range(3)
        ).lower()
        body = " ".join(
            comparison_table.cell(row, col).text
            for row in range(len(comparison_table.rows))
            for col in range(3)
        ).lower()
        assert "next.js" in header or "next" in header
        assert "drupal" in header
        assert "flutter" not in body
        assert "react native" not in body


def test_table_slides_keep_titles_after_footnote_clear():
    with tempfile.TemporaryDirectory() as tmp:
        output = str(Path(tmp) / "proposal.pptx")
        render_pptx_from_template(
            {
                "customer": "Acme Corp",
                "technicalApproach": "API Gateway | GraphQL federation layer",
                "sources": {"website": {"url": "https://acme.example", "pages_crawled": 3}},
                "understandingOfRequirements": "Corporate website redesign and CMS.",
            },
            output,
            "rfp-test-titles",
        )
        prs = Presentation(output)
        for index, expected in (
            (13, "Architecture Considerations"),
            (14, "Web Technology Platform Comparison"),
        ):
            titles = [
                shape.text.strip()
                for shape in prs.slides[index].shapes
                if hasattr(shape, "text_frame")
                and shape.top < 900_000
                and shape.text.strip()
            ]
            assert expected in titles


def test_render_clears_table_footnote():
    with tempfile.TemporaryDirectory() as tmp:
        output = str(Path(tmp) / "proposal.pptx")
        render_pptx_from_template(
            {
                "customer": "Acme Corp",
                "technicalApproach": "API Gateway | GraphQL federation layer",
            },
            output,
            "rfp-test-003",
        )

        slide = Presentation(output).slides[13]
        footnotes = [
            shape.text
            for shape in slide.shapes
            if hasattr(shape, "text_frame")
            and shape.top > 4_000_000
            and shape.text.strip()
            and shape.text.strip() != "‹#›"
        ]
        assert footnotes == []
        table_shapes = [
            shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE
        ]
        assert table_shapes
        table = table_shapes[0].table
        assert "API Gateway" in table.cell(1, 0).text


def test_fixed_slide_indices_cover_why_ttn_and_success_sections():
    assert WHY_TTN_SLIDE_INDEX in FIXED_SLIDE_INDICES
    assert SUCCESS_STORY_SLIDE_INDICES <= FIXED_SLIDE_INDICES
    assert 9 not in FIXED_SLIDE_INDICES
    assert 20 not in FIXED_SLIDE_INDICES
