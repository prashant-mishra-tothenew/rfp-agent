import tempfile
from pathlib import Path

from pptx import Presentation

from app.documents.parser import parse_document, parse_pptx


def _create_sample_pptx(path: Path) -> None:
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "Sample RFP Deck"
    title.placeholders[1].text = "Customer requirements overview"

    content = prs.slides.add_slide(prs.slide_layouts[1])
    content.shapes.title.text = "Requirements"
    content.placeholders[1].text = "Support 50,000+ SKUs\nMulti-currency checkout"
    prs.save(str(path))


def test_parse_pptx_extracts_slide_text():
    with tempfile.TemporaryDirectory() as tmp:
        pptx_path = Path(tmp) / "sample.pptx"
        _create_sample_pptx(pptx_path)

        result = parse_pptx(str(pptx_path))

    assert result["metadata"]["slide_count"] == 2
    assert "50,000+ SKUs" in result["text"]
    assert len(result["slides"]) == 2


def test_parse_document_routes_pptx():
    with tempfile.TemporaryDirectory() as tmp:
        pptx_path = Path(tmp) / "sample.pptx"
        _create_sample_pptx(pptx_path)

        result = parse_document(str(pptx_path))

    assert result["format"] == "pptx"
    assert result["filename"] == "sample.pptx"
