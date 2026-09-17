#!/usr/bin/env python3
"""Prepare proposal-template.pptx from a source TTN deck.

Removes the Racing Queensland client logo images while keeping the TTN logo.

Usage:
    python ingestion/scripts/prepare_proposal_template.py
    python ingestion/scripts/prepare_proposal_template.py \\
        --source "templates/TTN Proposal for Mobile App Development _ Racing Queensland v1.0 (2).pptx"
"""

import argparse
import hashlib
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "ai-service"))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = (
    REPO_ROOT / "templates" / "TTN Proposal for Mobile App Development .pptx"
)
DEFAULT_OUTPUT = REPO_ROOT / "templates" / "proposal-template.pptx"

# Image hashes for Racing Queensland branding to strip from the deck.
RACING_LOGO_HASHES = frozenset(
    {
        "d39af484c9c9",  # header client logo (top-right, left of TTN logo)
        "c93d98615591",  # cover / closing hero image
    }
)


def _image_hash(shape) -> str | None:
    try:
        return hashlib.md5(shape.image.blob).hexdigest()[:12]
    except Exception:
        return None


def _remove_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)


def _collect_removable_shapes(container) -> list:
    removable: list = []
    for shape in container.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            if _image_hash(shape) in RACING_LOGO_HASHES:
                removable.append(shape)
        elif hasattr(shape, "shapes"):
            for child in shape.shapes:
                if child.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    if _image_hash(child) in RACING_LOGO_HASHES:
                        removable.append(child)
    return removable


def remove_racing_logos(prs: Presentation) -> int:
    removed = 0
    containers = []
    for master in prs.slide_masters:
        containers.append(master)
        containers.extend(master.slide_layouts)
    containers.extend(prs.slides)

    for container in containers:
        for shape in _collect_removable_shapes(container):
            _remove_shape(shape)
            removed += 1

    return removed


def prepare_template(source: Path, output: Path) -> int:
    if not source.exists():
        raise FileNotFoundError(f"Source template not found: {source}")

    output.parent.mkdir(parents=True, exist_ok=True)
    staging = output.with_suffix(".staging.pptx")
    shutil.copy2(source, staging)

    prs = Presentation(str(staging))
    removed = remove_racing_logos(prs)
    prs.save(str(output))
    staging.unlink(missing_ok=True)
    return removed


def main() -> None:
    parser = argparse.ArgumentParser(description="Prepare proposal-template.pptx")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    removed = prepare_template(args.source, args.output)
    print(f"Prepared {args.output} ({args.output.stat().st_size} bytes)")
    print(f"Removed {removed} Racing Queensland image(s)")


if __name__ == "__main__":
    main()
